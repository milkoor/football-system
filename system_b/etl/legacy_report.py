"""舊版相容報表匯出模組（模板填充方式）。

以舊版 Report Excel 為模板，掃描 Location 標題列提取聯賽代碼，
用新系統的 decision_results 訊號填入對應儲存格。
模板中不存在於新系統的聯賽保持空白。
"""

from __future__ import annotations

import io
import logging
import re
from copy import copy

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

# Location 標題的正則：提取聯賽代碼（如 AUS1）
_LOC_RE = re.compile(r"^Location:\s*(\w+)")

# Signal cell offsets relative to the Location header row.
# Each group block = 1 header row + 5 zone rows.
# Group 1 data rows: offset +2 to +6  (header at +1)
# Group 2 data rows: offset +8 to +12 (header at +7)
_GROUP_OFFSETS = [
    (2, 6),   # group 1: rows header_row+2 .. header_row+6
    (8, 12),  # group 2: rows header_row+8 .. header_row+12
]

# Column indices for signal cells
_RT_HOME_COL = 3   # C
_RT_AWAY_COL = 6   # F
_EARLY_HOME_COL = 9   # I
_EARLY_AWAY_COL = 12  # L


def fill_template_report(
    template_bytes: bytes,
    decisions: list[dict],
    leagues_by_code: dict[str, int],
    group_ids: list[int],
) -> bytes:
    """以舊版 Report 為模板，填入新系統的訊號。

    Args:
        template_bytes: 舊版 Report Excel 的原始 bytes
        decisions: decision_results 列表（已篩選 play_type）
        leagues_by_code: league_code → league_id 對照表
        group_ids: 要填入的全域分組 ID 列表（最多 2 個，依序對應模板中的第一段/第二段）

    Returns:
        填充後的 Excel 檔案 bytes
    """
    wb = load_workbook(io.BytesIO(template_bytes))

    # Build lookup: (league_id, global_group_id, timing) -> decision
    idx: dict[tuple[int, int, str], dict] = {}
    for d in decisions:
        key = (d["league_id"], d.get("global_group_id", 0), d["timing"])
        idx[key] = d

    for ws in wb.worksheets:
        # Scan column A for Location headers
        for row in range(1, ws.max_row + 1):
            cell_val = ws.cell(row=row, column=1).value
            if not cell_val or not isinstance(cell_val, str):
                continue
            m = _LOC_RE.match(cell_val.strip())
            if not m:
                continue

            league_code = m.group(1)
            league_id = leagues_by_code.get(league_code)

            # For each group slot in the template
            for g_idx, (start_off, end_off) in enumerate(_GROUP_OFFSETS):
                if g_idx >= len(group_ids):
                    # No group assigned to this slot — clear cells
                    _clear_signal_cells(ws, row, start_off, end_off)
                    continue

                gid = group_ids[g_idx]

                # Get decisions for RT and Early
                d_rt = idx.get((league_id, gid, "RT")) if league_id else None
                d_early = idx.get((league_id, gid, "Early")) if league_id else None

                for zone_i in range(5):
                    data_row = row + start_off + zone_i

                    # RT signals
                    rt_home = ""
                    rt_away = ""
                    if d_rt:
                        rt_home = d_rt["home_signals"][zone_i] if zone_i < len(d_rt["home_signals"]) else ""
                        rt_away = d_rt["away_signals"][zone_i] if zone_i < len(d_rt["away_signals"]) else ""
                    ws.cell(row=data_row, column=_RT_HOME_COL).value = rt_home or None
                    ws.cell(row=data_row, column=_RT_AWAY_COL).value = rt_away or None

                    # Early signals
                    early_home = ""
                    early_away = ""
                    if d_early:
                        early_home = d_early["home_signals"][zone_i] if zone_i < len(d_early["home_signals"]) else ""
                        early_away = d_early["away_signals"][zone_i] if zone_i < len(d_early["away_signals"]) else ""
                    ws.cell(row=data_row, column=_EARLY_HOME_COL).value = early_home or None
                    ws.cell(row=data_row, column=_EARLY_AWAY_COL).value = early_away or None

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _clear_signal_cells(ws, header_row: int, start_off: int, end_off: int):
    """清空一個分組區塊的訊號儲存格。"""
    for zone_i in range(5):
        data_row = header_row + start_off + zone_i
        for col in (_RT_HOME_COL, _RT_AWAY_COL, _EARLY_HOME_COL, _EARLY_AWAY_COL):
            ws.cell(row=data_row, column=col).value = None
